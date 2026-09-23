"""Чтение и валидация рыночных данных заказчика (Excel).

Парсер не привязан к адресам ячеек: он находит блоки по заголовкам («Автопогрузчик…»,
«Электропогрузчик…»), строку периодов — по шаблону «N кв.ГГГГ», а сегменты — по подписи
грузоподъёмности. Поэтому новый файл с дополнительными кварталами (столбцами) или сегментами
читается без изменения кода — это основа автоматического пересчёта прогноза.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd

PERIOD_RE = re.compile(r"(\d)\s*кв\.?\s*(\d{4})", re.IGNORECASE)
NUM_RE = re.compile(r"\d+(?:[.,]\d+)?")

BLOCK_ORDER = {"ice": 0, "electric": 1}

BLOCK_PATTERNS = {
    "ice": re.compile(r"автопогрузчик", re.IGNORECASE),
    "electric": re.compile(r"электропогрузчик", re.IGNORECASE),
}


def parse_period(label: Any) -> pd.Period | None:
    """``'1 кв.2023'`` → ``Period('2023Q1')``; иначе ``None``."""
    if label is None:
        return None
    m = PERIOD_RE.search(str(label))
    if not m:
        return None
    return pd.Period(f"{m.group(2)}Q{m.group(1)}", freq="Q")


def parse_capacity(label: str) -> tuple[float, float]:
    """Извлекает диапазон грузоподъёмности из подписи: ``'3,8-5,0 т.'`` → ``(3.8, 5.0)``."""
    nums = [float(x.replace(",", ".")) for x in NUM_RE.findall(str(label))]
    if not nums:
        raise ValueError(f"Не удалось распознать грузоподъёмность в подписи {label!r}")
    return (nums[0], nums[-1]) if len(nums) > 1 else (nums[0], nums[0])


@dataclass
class DataQualityReport:
    """Результаты проверок качества данных — попадают в отчёт и в лог."""

    checks: list[dict[str, Any]] = field(default_factory=list)

    def add(self, name: str, status: str, detail: str) -> None:
        self.checks.append({"проверка": name, "статус": status, "детали": detail})

    @property
    def ok(self) -> bool:
        return all(c["статус"] != "ОШИБКА" for c in self.checks)

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame(self.checks)


def _segment_code(block: str, lo: float, hi: float, segments_cfg: dict[str, Any]) -> tuple[str, str, bool]:
    for seg in segments_cfg.get("segments", {}).get(block, []):
        if abs(seg["capacity_min"] - lo) < 1e-9:
            return seg["code"], seg["label"], bool(seg.get("silant", False))
    # новый, ранее не встречавшийся сегмент — создаём код автоматически
    prefix = "ice" if block == "ice" else "el"
    fmt = lambda x: f"{x:g}"  # noqa: E731
    return f"{prefix}_{fmt(lo)}_{fmt(hi)}", f"{fmt(lo)}–{fmt(hi)} т".replace(".", ","), False


def read_market_excel(path: str | Path, segments_cfg: dict[str, Any]) -> tuple[pd.DataFrame, DataQualityReport]:
    """Читает Excel с объёмами рынка и возвращает длинную таблицу + отчёт о качестве.

    Returns
    -------
    df : DataFrame со столбцами ``period, block, segment, label, capacity_min, capacity_max,
         silant, units``
    dq : DataQualityReport
    """
    path = Path(path)
    dq = DataQualityReport()
    wb_values = openpyxl.load_workbook(path, data_only=True)
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    ws = wb_formulas.worksheets[0]
    ws_values = wb_values.worksheets[0]

    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    rows_values = [[c for c in r] for r in ws_values.iter_rows(values_only=True)]
    records: list[dict[str, Any]] = []
    block: str | None = None
    periods: dict[int, pd.Period] = {}
    totals_in_file: dict[tuple[str, pd.Period], Any] = {}

    for i, row in enumerate(rows):
        text = " ".join(str(v) for v in row if v is not None)
        for b, pat in BLOCK_PATTERNS.items():
            if pat.search(text) and not any(isinstance(v, (int, float)) for v in row):
                block, periods = b, {}
        if block is None:
            continue
        found = {j: parse_period(v) for j, v in enumerate(row)}
        found = {j: p for j, p in found.items() if p is not None}
        if len(found) >= 2:
            periods = found
            continue
        if not periods or row[0] is None:
            continue
        label = str(row[0]).strip()
        if label.lower().startswith("итого"):
            for j, p in periods.items():
                totals_in_file[(block, p)] = (row[j], rows_values[i][j] if i < len(rows_values) else None)
            continue
        try:
            lo, hi = parse_capacity(label)
        except ValueError:
            continue
        code, nice_label, silant = _segment_code(block, lo, hi, segments_cfg)
        for j, p in periods.items():
            v = row[j]
            records.append(
                {
                    "period": p,
                    "block": block,
                    "segment": code,
                    "label": nice_label,
                    "source_label": label,
                    "power": row[1] if len(row) > 1 else None,
                    "capacity_min": lo,
                    "capacity_max": hi,
                    "silant": silant,
                    "units": v,
                }
            )

    df = pd.DataFrame.from_records(records)
    if df.empty:
        dq.add("структура файла", "ОШИБКА", "Не найдено ни одного блока с данными")
        return df, dq

    # --- проверки качества ----------------------------------------------------------------------
    non_numeric = df[pd.to_numeric(df["units"], errors="coerce").isna()]
    if len(non_numeric):
        dq.add("числовые значения", "ОШИБКА", f"{len(non_numeric)} нечисловых/пустых ячеек — заменены на 0")
    else:
        dq.add("числовые значения", "OK", "все значения числовые, пропусков нет")
    df["units"] = pd.to_numeric(df["units"], errors="coerce").fillna(0.0).astype(float)

    neg = int((df["units"] < 0).sum())
    dq.add("неотрицательность", "OK" if neg == 0 else "ОШИБКА", f"отрицательных значений: {neg}")

    per = sorted(df["period"].unique())
    expected = pd.period_range(per[0], per[-1], freq="Q")
    gaps = sorted(set(expected) - set(per))
    dq.add(
        "непрерывность ряда",
        "OK" if not gaps else "ОШИБКА",
        f"{len(per)} кварталов: {per[0]}–{per[-1]}" + (f"; пропуски: {gaps}" if gaps else ""),
    )

    dup = int(df.duplicated(["period", "segment"]).sum())
    dq.add("дубликаты", "OK" if dup == 0 else "ОШИБКА", f"дублирующихся записей: {dup}")

    for b in df["block"].unique():
        n = df.loc[df.block == b, "segment"].nunique()
        dq.add(f"сегменты: {b}", "INFO", f"{n} сегментов")

    mult10 = float((df["units"] % 10 == 0).mean())
    dq.add(
        "округление",
        "ВНИМАНИЕ" if mult10 > 0.95 else "INFO",
        f"{mult10:.0%} значений кратны 10 → данные округлены; для малых сегментов (<100 шт./кв.) "
        "ошибка округления ±5 шт. сопоставима с сигналом",
    )

    zero_share = df.groupby("segment")["units"].apply(lambda s: float((s == 0).mean()))
    sparse = zero_share[zero_share >= 0.3]
    if len(sparse):
        dq.add(
            "разреженные ряды",
            "ВНИМАНИЕ",
            "нулевые значения в ≥30% кварталов: " + ", ".join(f"{k} ({v:.0%})" for k, v in sparse.items())
            + " — прогнозируются через структуру блока, а не отдельной моделью",
        )

    # сверка строк «Итого» (формулы в файле) с суммой сегментов
    calc = df.groupby(["block", "period"])["units"].sum()
    mismatches = []
    for (b, p), (formula, cached) in totals_in_file.items():
        if isinstance(formula, (int, float)):
            if abs(float(formula) - calc.get((b, p), np.nan)) > 0.5:
                mismatches.append(f"{b} {p}: в файле {formula}, сумма {calc.get((b, p))}")
        elif isinstance(cached, (int, float)) and abs(float(cached) - calc.get((b, p), np.nan)) > 0.5:
            mismatches.append(f"{b} {p}: кэш {cached}, сумма {calc.get((b, p))}")
    dq.add(
        "строки «Итого»",
        "OK" if not mismatches else "ВНИМАНИЕ",
        "итоги пересчитаны из сегментов (в файле — формулы SUM)" if not mismatches else "; ".join(mismatches),
    )
    df = (
        df.assign(_b=df["block"].map(BLOCK_ORDER))
        .sort_values(["_b", "capacity_min", "period"])
        .drop(columns="_b")
        .reset_index(drop=True)
    )
    return df, dq


def to_wide(df: pd.DataFrame) -> pd.DataFrame:
    """Широкая таблица: индекс — квартал, столбцы — сегменты + итоги блоков и рынка."""
    wide = df.pivot_table(index="period", columns="segment", values="units", aggfunc="sum")
    seg = df.drop_duplicates("segment").assign(_b=lambda d: d["block"].map(BLOCK_ORDER))
    order = seg.sort_values(["_b", "capacity_min"])["segment"].tolist()
    wide = wide[order]
    for b in ["ice", "electric"]:
        cols = df.loc[df.block == b, "segment"].unique()
        wide[f"total_{b}"] = wide[list(cols)].sum(axis=1)
    wide["total_market"] = wide["total_ice"] + wide["total_electric"]
    wide["electric_share"] = wide["total_electric"] / wide["total_market"]
    wide.index = pd.PeriodIndex(wide.index, freq="Q")
    return wide

"""Краткая презентация (PowerPoint, 16:9) с выводами и рекомендациями для планирования."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .text import key_findings, recommendations

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x89, 0x85)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
SOFT = RGBColor(0xF3, 0xF2, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)

SHORT_LIMITS = [
    ("Короткая история", "13 кварталов: оценки стабилизированы априорными распределениями, интервалы намеренно широкие."),
    ("Нет цен, брендов, регионов", "нельзя оценить ценовую эластичность и конкуренцию; доля СИЛАНТ задаётся экспертно."),
    ("8 сегментов ДВС вместо 9", "в файле нет девятого сегмента из описания кейса; прогноз — по фактической структуре."),
    ("Округление до 10 шт.", "малые сегменты (6–25 т) шумные; прогнозируются через структуру блока."),
    ("Структурный сдвиг 2025Q2", "причины (конец цикла замещения, утильсбор, кредитование) разделить нельзя."),
    ("Условный бэктест", "факторы в проверке — фактические; неопределённость будущих факторов — через сценарии."),
]


def _text(slide, x, y, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FONT
    return tb


def _bullets(slide, x, y, w, h, items, size=15, color=INK, bold_head=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        if isinstance(it, tuple):
            r = p.add_run()
            r.text = it[0] + ". "
            r.font.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(size), ACCENT, FONT
            r2 = p.add_run()
            r2.text = it[1]
            r2.font.size, r2.font.color.rgb, r2.font.name = Pt(size), color, FONT
        else:
            r = p.add_run()
            r.text = "• " + it
            r.font.size, r.font.color.rgb, r.font.name = Pt(size), color, FONT
    return tb


def _header(slide, title, subtitle=None, num=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8), title, size=28, bold=True)
    if subtitle:
        _text(slide, Inches(0.5), Inches(0.98), Inches(12.3), Inches(0.5), subtitle, size=15, color=INK2)
    if num:
        _text(slide, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.35), str(num), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    _text(slide, Inches(0.5), Inches(7.0), Inches(8), Inches(0.35), "Forkcast · рынок вилочных погрузчиков РФ · ЧЗСА / СИЛАНТ",
          size=10, color=MUTED)


def _kpi(slide, x, y, w, h, value, label):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = SOFT
    box.line.fill.background()
    box.adjustments[0] = 0.08
    _text(slide, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.7), value, size=30, bold=True, color=ACCENT)
    _text(slide, x + Inches(0.2), y + Inches(0.9), w - Inches(0.4), h - Inches(1.0), label, size=13, color=INK2)


def _pic(slide, path, x, y, w=None, h=None):
    if path and Path(path).exists():
        return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


def _style_chart(chart, colors):
    chart.font.size = Pt(11)
    chart.font.name = FONT
    for s, c in zip(chart.plots[0].series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c


def build_presentation(res, figs: dict[str, Path], path: Path) -> Path:
    k = key_findings(res)
    recs = recommendations(res)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    n = 0

    def new(notes: str | None = None):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # 1. Титул
    s = new("Цель: формализованная модель прогноза рынка по сегментам для производственного и продуктового планирования.")
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ACCENT
    bg.line.fill.background()
    _text(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.6), "FORKCAST", size=20, bold=True, color=RGBColor(0xB7, 0xD3, 0xF6))
    _text(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.8),
          "Рынок вилочных погрузчиков РФ: анализ сегментов и сценарный прогноз", size=40, bold=True, color=WHITE)
    _text(s, Inches(0.8), Inches(4.2), Inches(11.5), Inches(0.8),
          f"Данные {res.facts['first_period']}–{k['last_period']} · прогноз до {k['fc_8q_end']} · для ЧЗСА (бренд СИЛАНТ)",
          size=18, color=WHITE)
    _text(s, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.8),
          ["Сергеев Антон Валентинович", "avsergeev1981@gmail.com"], size=14, color=RGBColor(0xDD, 0xE8, 0xF7))

    # 2. Главное
    s = new()
    _header(s, "Главное", "Рынок сжался, дно пройдено, восстановление медленное; электрификация продолжается", n)
    kw = Inches(2.95)
    for i, (v, lab) in enumerate([
        (k["market_y1"], f"рынок {k['y1']} г., шт. ({k['market_d']} к {k['y0']})"),
        (k["w4_p50"], f"прогноз на 4 кв. ({k['fc_start']}–{k['fc_4q_end']}), P10–P90: {k['w4_p10']}–{k['w4_p90']}"),
        (k["share_8q"], f"доля электро к {k['fc_8q_end']} (сейчас {k['last_share']})"),
        (k["beta_kr"], "изменение продаж на +1 п.п. ключевой ставки (лаг 1–3 кв.)"),
    ]):
        _kpi(s, Inches(0.5) + i * (kw + Inches(0.15)), Inches(1.6), kw, Inches(1.9), v, lab)
    _bullets(s, Inches(0.5), Inches(3.8), Inches(12.3), Inches(3.0), [
        f"Падение {k['y1']} г. ({k['dec_total']} лог-п.): ставка {k['dec_kr']}, структурный сдвиг 2025Q2 {k['dec_reg']}, инвестиции {k['dec_inv']}.",
        f"ДВС {k['ice_d']}, электро {k['el_d']} в {k['y1']} г. — ДВС сокращается быстрее рынка, доля электро {k['share_y1']}.",
        f"Сценарии почти не расходятся на 4 кв. ({k['sc4_min']}–{k['sc4_max']}), но расходятся на 8 кв. ({k['sc8_min']}–{k['sc8_max']}).",
        f"Модель точнее эталонов в бэктесте: MASE {k['bt_mase']} против {k['bt_mase_sn']} у сезонного наивного прогноза.",
    ], size=16)

    # 3. История рынка (нативная диаграмма)
    s = new()
    _header(s, f"Рынок {k['market_d']} в {k['y1']} г., ДВС {k['ice_d']}", "Продажи по кварталам, шт.", n)
    cd = CategoryChartData()
    cd.categories = [f"{p.quarter}кв{str(p.year)[2:]}" for p in res.wide.index]
    cd.add_series("ДВС", [float(v) for v in res.wide["total_ice"]])
    cd.add_series("Электро", [float(v) for v in res.wide["total_electric"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(1.5), Inches(8.3), Inches(5.3), cd)
    ch = gf.chart
    _style_chart(ch, [ORANGE, BLUE])
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.plots[0].gap_width = 60
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE5, 0xE1)
    ch.value_axis.format.line.fill.background()
    ch.value_axis.tick_labels.font.color.rgb = INK2
    ch.category_axis.tick_labels.font.color.rgb = INK2
    _bullets(s, Inches(9.1), Inches(1.6), Inches(3.9), Inches(5.2), [
        f"{k['y1']} г.: {k['market_y1']} шт. ({k['market_d']}).",
        f"ДВС {k['ice_d']}, электро {k['el_d']}.",
        f"{k['last_period']}: {k['last_total']} шт. ({k['last_yoy']} г/г) — минимум ряда.",
        f"Доля электро: {k['share_2023']} (2023) → {k['share_y1']} ({k['y1']}).",
        f"Сезонность значима: 4 кв. {k['q4_eff']}, 1 кв. {k['q1_eff']} к среднегодовому.",
        "Растут только электро 3–3,5 т и 3,8–5 т.",
    ], size=14)

    # 4. Драйверы
    s = new()
    _header(s, "Что двигает рынок", "Вклад факторов в изменение рынка и оценённые эластичности", n)
    _pic(s, figs.get("decomposition"), Inches(0.5), Inches(1.5), w=Inches(7.6))
    _bullets(s, Inches(8.4), Inches(1.6), Inches(4.6), Inches(5.2), [
        ("Ключевая ставка", f"{k['beta_kr']} продаж на +1 п.п.; лаг 1–3 квартала (кредит, лизинг)."),
        ("Инвестиции в ОК", f"эластичность {k['beta_inv']}: спрос на технику — «усиленная» производная капвложений."),
        ("Структурный сдвиг 2025Q2", f"{k['beta_regime']} к уровню: конец цикла замещения парка, утильсбор, неценовые условия кредитования."),
        ("Отклонены", "курс USD/RUB и ВВП — не улучшают точность вне выборки (LOO)."),
    ], size=14)

    # 5. Методика
    s = new("Байесовская иерархическая модель: априорные распределения стабилизируют оценки на 13 точках; апостериорные выборки дают неопределённость параметров.")
    _header(s, "Методика: иерархическая байесовская факторная модель", "Прогнозы согласованы: сегменты → блоки → рынок", n)
    boxes = [
        ("Внешние факторы", "Ставка ЦБ (лаг 1–3), инвестиции в ОК, режим 2025Q2 · 5 сценариев (ЦБ, Минэк, консенсус, стресс, оптимистичный)"),
        ("Рынок в целом", "log-уровень продаж = сезонность + факторы; байесовская регрессия, сэмплер Гиббса"),
        ("Доля электро", "S-кривая с насыщением (потолок 85%) + циклическая реакция на ставку"),
        ("14 сегментов", "композиционная модель структуры внутри ДВС и электро; сумма = блок = рынок"),
        ("План СИЛАНТ", "коридор P10–P90 и объём выпуска по правилу newsvendor для сегментов 1–5 т"),
    ]
    for i, (h_, b_) in enumerate(boxes):
        y = Inches(1.6) + i * Inches(1.05)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), y, Inches(3.2), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = ACCENT if i in (1, 2) else SOFT
        box.line.fill.background()
        _text(s, Inches(0.6), y, Inches(3.0), Inches(0.85), h_, size=16, bold=True,
              color=WHITE if i in (1, 2) else ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        _text(s, Inches(3.9), y, Inches(5.0), Inches(0.85), b_, size=13, color=INK2, anchor=MSO_ANCHOR.MIDDLE)
    _bullets(s, Inches(9.2), Inches(1.6), Inches(3.9), Inches(5.2), [
        f"Неопределённость: параметры + шум + путь факторов + выбор сценария; {res.manifest['n_draws']} симуляций на сценарий.",
        "Наукаст: уже известные ставка и инвестиции заменяют сценарные значения.",
        f"Бэктест (rolling origin): MASE {k['bt_mase']}, попадание в 80%-интервал {k['bt_cov80']}.",
        "Отбор факторов — по кросс-валидации LOO, а не по подгонке.",
    ], size=13)

    # 6. Прогноз рынка
    s = new()
    _header(s, f"Прогноз рынка: {k['w4_p50']} шт. за 4 квартала", f"Взвешенный сценарий; 80%-интервал {k['w4_p10']}–{k['w4_p90']}; на 8 кв. — {k['w8_p50']} шт.", n)
    _pic(s, figs.get("fan_market"), Inches(0.5), Inches(1.5), h=Inches(5.4))

    # 7. Сценарии (нативная диаграмма)
    s = new()
    _header(s, "Сценарии расходятся во втором году", "Рынок за 4 и 8 кварталов по сценариям, медиана, шт.", n)
    st = res.tables["scenario_totals"]
    stm = st[st.series == "total_market"]
    cd = CategoryChartData()
    names = [("Взвешенный" if r.scenario == "weighted" else res.scenario(r.scenario).title.split("(")[0].split(":")[0].strip())
             for r in stm.itertuples()]
    cd.categories = names
    cd.add_series("4 квартала", [float(v) for v in stm["4q_p50"]])
    cd.add_series("8 кварталов", [float(v) for v in stm["8q_p50"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.5), Inches(7.6), Inches(5.3), cd)
    ch = gf.chart
    _style_chart(ch, [BLUE, ACCENT])
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.number_format = '# ##0'
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    pl.data_labels.font.size = Pt(10)
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE5, 0xE1)
    ch.category_axis.reverse_order = True
    ch.value_axis.tick_labels.number_format = '# ##0'
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.font.size = Pt(9)
    items = []
    for sc in res.scenarios:
        items.append((sc.title.split("(")[0].strip() + f" ({sc.weight:.0%})",
                      f"ставка {sc.path['key_rate'].iloc[2]:.1f}% в 4кв26 → {sc.path['key_rate'].iloc[6]:.1f}% в 4кв27; "
                      f"инвестиции 2027 ≈ {sc.path['investment_yoy'].iloc[3:7].mean():+.1f}% г/г"))
    _bullets(s, Inches(8.3), Inches(1.5), Inches(4.8), Inches(5.4), items, size=12)

    # 8. Доля электро
    s = new()
    _header(s, f"Доля электропогрузчиков: {k['share_8q']} к {k['fc_8q_end']}",
            f"S-кривая замещения; 80%-интервал {k['share_8q_lo']}–{k['share_8q_hi']}; слабо зависит от сценария", n)
    _pic(s, figs.get("fan_share"), Inches(0.5), Inches(1.5), h=Inches(5.4))

    # 9. Сегменты СИЛАНТ
    s = new()
    _header(s, "Сегменты линейки СИЛАНТ (1–5 т)", "Медиана и 80%-интервал; коридор плана — в Excel (лист «План_СИЛАНТ»)", n)
    _pic(s, figs.get("silant"), Inches(0.5), Inches(1.45), w=Inches(12.3))

    # 10. Неопределённость и точность
    s = new()
    _header(s, "Неопределённость оценена и проверена на истории", "Бэктест с точками отсечения 2024Q4–2025Q4; сравнение с эталонными моделями", n)
    _pic(s, figs.get("backtest"), Inches(0.5), Inches(1.5), w=Inches(8.4))
    _pic(s, figs.get("uncertainty"), Inches(9.0), Inches(1.6), w=Inches(4.1))
    _bullets(s, Inches(9.0), Inches(4.0), Inches(4.1), Inches(2.8), [
        f"MASE {k['bt_mase']} vs {k['bt_mase_sn']} (SNaive).",
        f"Факт внутри 80%-интервала: {k['bt_cov80']}.",
        "Ближний горизонт — шум рынка; дальний — неопределённость факторов → сценарии.",
    ], size=12)

    # 11. Лестница ставок
    s = new()
    _header(s, "«Лестница ставок»: при какой ставке рынок вернётся к росту",
            f"Рынок {k['ladder_year']} г. при постоянной ключевой ставке; уровень {k['y1']} г. — при ставке ≤ {k['breakeven']}%", n)
    _pic(s, figs.get("ladder"), Inches(0.5), Inches(1.6), w=Inches(8.5))
    _bullets(s, Inches(9.3), Inches(1.7), Inches(3.8), Inches(5), [
        "Ставка действует с лагом 1–3 квартала: решения ЦБ 2026 г. определяют рынок первой половины 2027 г.",
        "Триггер пересмотра плана: отклонение ставки от сценария более чем на 1 п.п.",
        "Следующее заседание ЦБ — 23.10.2026.",
    ], size=13)

    # 12. Рекомендации
    s = new()
    _header(s, "Рекомендации для производственного и продуктового плана", None, n)
    _bullets(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.8), recs, size=13)

    # 13. Ограничения и развитие
    s = new()
    _header(s, "Ограничения и развитие модели", None, n)
    _bullets(s, Inches(0.5), Inches(1.2), Inches(7.6), Inches(5.8), SHORT_LIMITS, size=13)
    _bullets(s, Inches(8.4), Inches(1.2), Inches(4.7), Inches(5.8), [
        ("Данные", "подключить ввод складов, лизинг спецтехники, импорт погрузчиков (таможня КНР), цены."),
        ("Модель", "доля СИЛАНТ по сегментам из CRM → план продаж, а не только рынка."),
        ("Процесс", "ежеквартальный пересчёт, мониторинг PIT, пересмотр весов сценариев."),
    ], size=13)

    # 14. Инструмент
    s = new()
    _header(s, "Инструмент: пересчёт одной командой", "Воспроизводимый код, тесты, архив прогнозов и мониторинг точности", n)
    _text(s, Inches(0.5), Inches(1.6), Inches(7.5), Inches(3.2), [
        "forkcast update --market новый_файл.xlsx",
        "forkcast fetch-macro",
        "forkcast run",
        "forkcast backtest",
    ], size=18, color=ACCENT)
    _bullets(s, Inches(7.8), Inches(1.6), Inches(5.3), Inches(5.2), [
        "Новый квартал или сегмент распознаётся автоматически — без правки кода.",
        "Сценарии и веса — в YAML-конфигурации, без программирования.",
        "Выходы: Excel для планировщиков, HTML-отчёт, интерактивный дашборд, эта презентация.",
        "Каждый прогноз архивируется; при поступлении факта считается его точность.",
        "Фиксированный seed и манифест с хэшами данных — результат воспроизводим до знака.",
    ], size=14)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path

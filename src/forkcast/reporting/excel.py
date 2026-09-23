"""Excel-выгрузка для планировщиков: все прогнозы, допущения, план и датасет в одном файле."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
HEADER_FONT = Font(bold=True, color="FFFFFF")

SERIES_RU = {
    "total_market": "Рынок, всего", "total_ice": "ДВС, всего", "total_electric": "Электро, всего",
    "electric_share": "Доля электро",
}


def _series_name(res, code: str) -> str:
    if code in SERIES_RU:
        return SERIES_RU[code]
    blk = "ДВС" if code.startswith("ice") else "Электро"
    return f"{blk} {res.labels.get(code, code)}"


def _style(ws, df: pd.DataFrame, pct_cols=(), int_cols=(), float_cols=()):
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        width = max(10, min(60, int(max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) * 1.1)))
        ws.column_dimensions[get_column_letter(j)].width = width
        fmt = "0.0%" if col in pct_cols else "# ##0" if col in int_cols else "0.000" if col in float_cols else None
        if fmt:
            for i in range(2, len(df) + 2):
                ws.cell(row=i, column=j).number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_excel(res, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    t = res.tables
    sheets: list[tuple[str, pd.DataFrame, dict]] = []

    # 1. Сводка
    st = t["scenario_totals"].copy()
    st["series"] = st["series"].map(lambda c: _series_name(res, c))
    fp = res.facts["fc_periods"]
    st = st.rename(columns={
        "scenario": "Код сценария", "title": "Сценарий", "weight": "Вес", "series": "Показатель",
        "4q_p10": f"4 кв. ({fp[0]}–{fp[3]}) P10", "4q_p50": "4 кв. P50", "4q_p90": "4 кв. P90",
        "8q_p10": f"8 кв. ({fp[0]}–{fp[-1]}) P10", "8q_p50": "8 кв. P50", "8q_p90": "8 кв. P90",
    })
    sheets.append(("Сводка", st, {"pct_cols": ["Вес"]}))

    # 2. Прогноз по сериям (все сценарии)
    sm = res.summary.copy()
    sm["period"] = sm["period"].astype(str)
    sm["Показатель"] = sm["series"].map(lambda c: _series_name(res, c))
    sm = sm.rename(columns={"scenario": "Сценарий", "period": "Квартал", "mean": "Среднее", "q500": "Медиана (P50)",
                            "q100": "P10", "q900": "P90", "q025": "P2.5", "q975": "P97.5", "q250": "P25", "q750": "P75",
                            "q050": "P5", "q950": "P95"})
    cols = ["Сценарий", "Показатель", "Квартал", "Медиана (P50)", "Среднее", "P2.5", "P5", "P10", "P25", "P75", "P90",
            "P95", "P97.5"]
    fc_units = sm[sm["series"] != "electric_share"][cols]
    fc_share = sm[sm["series"] == "electric_share"][cols]
    sheets.append(("Прогноз_шт", fc_units, {"int_cols": cols[3:]}))
    sheets.append(("Прогноз_доля_электро", fc_share, {"pct_cols": cols[3:]}))

    # 3. Широкая таблица медиан взвешенного прогноза: сегменты × кварталы
    wsum = res.weighted.summary()
    piv = wsum[wsum.series != "electric_share"].pivot(index="series", columns="period", values="q500")
    piv.columns = [str(c) for c in piv.columns]
    order = [c for c in res.wide.columns if c in piv.index]
    piv = piv.loc[order]
    hist = res.wide[order].T
    hist.columns = [str(c) for c in hist.columns]
    both = pd.concat([hist, piv], axis=1)
    both.insert(0, "Показатель", [_series_name(res, c) for c in both.index])
    sheets.append(("Факт+прогноз_P50", both.reset_index(drop=True), {"int_cols": list(both.columns[1:])}))

    # 4. План СИЛАНТ
    for key, name in (("plan_short", "План_СИЛАНТ_4кв"), ("plan_long", "План_СИЛАНТ_8кв")):
        pl = t[key].copy()
        pl["period"] = pl["period"].astype(str)
        pl["block"] = pl["block"].map({"ice": "ДВС", "electric": "Электро"})
        pl = pl.rename(columns={
            "block": "Блок", "segment": "Код", "label": "Сегмент", "period": "Квартал",
            "market_p10": "Рынок P10", "market_p50": "Рынок P50", "market_p90": "Рынок P90",
            "silant_share": "Целевая доля СИЛАНТ", "silant_demand_p50": "Спрос СИЛАНТ P50",
            "plan_newsvendor": "Рекомендуемый выпуск (newsvendor)", "critical_ratio": "Критический квантиль",
            "P(излишек)": "Вероятность излишка",
        })
        sheets.append((name, pl, {"int_cols": ["Рынок P10", "Рынок P50", "Рынок P90", "Спрос СИЛАНТ P50",
                                              "Рекомендуемый выпуск (newsvendor)"],
                                  "pct_cols": ["Целевая доля СИЛАНТ", "Критический квантиль", "Вероятность излишка"]}))

    # 5. Сценарии и допущения
    rows = []
    for sc in res.scenarios:
        for p, r in sc.path.iterrows():
            rows.append({"Сценарий": sc.title, "Вес": sc.weight, "Квартал": str(p),
                         "Ключевая ставка, %": r["key_rate"], "Инвестиции в ОК, % г/г": r["investment_yoy"],
                         "USD/RUB": r["usd_rub"], "ВВП, % г/г": r["gdp_yoy"], "Допущения": sc.assumptions})
    sheets.append(("Сценарии_допущения", pd.DataFrame(rows), {"pct_cols": ["Вес"]}))

    # 6. Модель и факторы
    sheets.append(("Коэф_рынок", t["coef_market"], {"float_cols": list(t["coef_market"].columns[1:])}))
    sheets.append(("Коэф_доля_электро", t["coef_share"], {"float_cols": list(t["coef_share"].columns[1:])}))
    sheets.append(("Отбор_факторов", t["spec_comparison"], {"float_cols": ["R2", "AIC", "LOO_RMSE"]}))
    sheets.append(("Факторы_long_list", t["factor_longlist"], {}))
    sheets.append(("Скрининг_корреляций", t["factor_screening"], {"float_cols": list(t["factor_screening"].columns[1:])}))
    dec = t["decomposition"].copy()
    sheets.append(("Декомпозиция", dec, {"float_cols": list(dec.columns[2:])}))
    sheets.append(("Лестница_ставок", t["key_rate_ladder"], {"int_cols": [c for c in t["key_rate_ladder"].columns if "_" in c and "share" not in c]}))
    ub = t["uncertainty_budget"].copy()
    ub["period"] = ub["period"].astype(str)
    sheets.append(("Бюджет_неопределенности", ub, {"pct_cols": list(ub.columns[1:4])}))
    for k, nm in (("prior_sensitivity", "Чувств_априори"), ("share_cap_sensitivity", "Чувств_потолок_доли")):
        if k in t:
            sheets.append((nm, t[k], {"float_cols": list(t[k].columns)}))

    # 7. Аналитика
    sheets.append(("Тренды_годовые", t["annual"], {"int_cols": [c for c in t["annual"].columns if c.isdigit()],
                                                   "pct_cols": [c for c in t["annual"].columns if c.startswith("Δ")]}))
    sheets.append(("Структура_сегментов", t["structure"], {"pct_cols": [c for c in t["structure"].columns if c.isdigit()]}))
    sheets.append(("Сезонность", t["seasonality"], {"pct_cols": [c for c in t["seasonality"].columns if c.startswith("Q")]}))

    # 8. Бэктест
    bs = t["backtest_scores"].copy()
    sheets.append(("Бэктест_метрики", bs, {"pct_cols": ["MAPE", "bias_%", "cover80", "cover95"], "float_cols": ["MASE"]}))
    bt = t["backtest"].copy()
    bt["origin"] = bt["origin"].astype(str)
    bt["period"] = bt["period"].astype(str)
    sheets.append(("Бэктест_детально", bt, {}))
    if "vintage_eval" in t and len(t["vintage_eval"]):
        sheets.append(("Мониторинг_точности", t["vintage_eval"], {}))

    # 9. Данные
    sheets.append(("Качество_данных", t["data_quality"], {}))
    ds = res.wide.join(res.features, how="left").copy()
    ds.index = ds.index.astype(str)
    sheets.append(("Датасет_модели", ds.reset_index(names="period"), {}))
    mc = res.macro.copy()
    mc.index = mc.index.astype(str)
    sheets.append(("Макро_квартал", mc.reset_index(names="period"), {}))
    src = yaml.safe_load(open(res.settings.external_dir / "sources.yaml", encoding="utf-8"))
    srows = []
    for key, v in {**src.get("indicators", {}), **src.get("forecasts", {})}.items():
        srows.append({"код": key, "показатель": v.get("title"), "источник": v.get("source", ""),
                      "ссылки": "\n".join(v.get("urls", [])), "статус_проверки": v.get("verification", ""),
                      "дата_выгрузки": src.get("retrieved_at"), "примечание": v.get("note", "")})
    sheets.append(("Источники", pd.DataFrame(srows), {}))

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df, fmt in sheets:
            df = df.replace([np.inf, -np.inf], np.nan)
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _style(xw.sheets[name[:31]], df, **fmt)
    return path

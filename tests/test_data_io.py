import openpyxl
import pandas as pd
import pytest

from forkcast.data_io import parse_capacity, parse_period, read_market_excel, to_wide


@pytest.mark.parametrize("label,expected", [("1 кв.2023", "2023Q1"), ("4 кв. 2025", "2025Q4"), ("итого", None)])
def test_parse_period(label, expected):
    p = parse_period(label)
    assert (str(p) if p is not None else None) == expected


@pytest.mark.parametrize("label,expected", [("1-1,8т.", (1.0, 1.8)), ("3,8-5,0 т.", (3.8, 5.0)),
                                            ("от 20 до 25 т.", (20.0, 25.0)), ("3-3,5", (3.0, 3.5))])
def test_parse_capacity(label, expected):
    assert parse_capacity(label) == expected


def test_reads_customer_file(market):
    long, dq, wide = market
    assert dq.ok
    assert long["segment"].nunique() == 14
    assert long[long.block == "ice"]["segment"].nunique() == 8
    assert len(wide) == 13
    assert str(wide.index[0]) == "2023Q1" and str(wide.index[-1]) == "2026Q1"
    # контрольные суммы из исходного файла
    assert wide.loc[pd.Period("2023Q1", "Q"), "total_ice"] == 13340
    assert wide.loc[pd.Period("2026Q1", "Q"), "total_electric"] == 10800
    assert (wide["total_market"] == wide["total_ice"] + wide["total_electric"]).all()


def test_new_quarter_is_detected(settings, tmp_path):
    """Новый столбец (квартал) распознаётся без изменения кода — основа автообновления."""
    wb = openpyxl.load_workbook(settings.market_file)
    ws = wb.worksheets[0]
    col = ws.max_column + 1
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=2).value
        head = ws.cell(row=row, column=3).value
        if isinstance(head, str) and "кв" in head:
            ws.cell(row=row, column=col, value="2 кв.2026")
        elif isinstance(ws.cell(row=row, column=3).value, (int, float)):
            ws.cell(row=row, column=col, value=ws.cell(row=row, column=col - 1).value)
        del v
    path = tmp_path / "new.xlsx"
    wb.save(path)
    long, dq = read_market_excel(path, settings.segments)
    w = to_wide(long)
    assert str(w.index[-1]) == "2026Q2"
    assert dq.ok
